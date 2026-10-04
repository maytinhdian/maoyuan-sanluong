import sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule, FormulaRule

NL = 9            # last seeded NHAP_LIEU row (table end)
N = 10004         # formula ranges over NHAP_LIEU reach this row
MN = 2004         # formula ranges over MUC_TIEU_THANG reach this row
LN = 10           # HIEN_THI line rows 5..10 (6 lines)
LR = 1000         # list ranges (lines, products)
HN = 10004        # formula ranges over HANG_LOI reach this row
SAMPLE = len(sys.argv) > 2 and sys.argv[2] == 'sample'   # bản mẫu cho app: thêm số giờ và nhiều dòng lỗi có ảnh
F0 = 'Calibri'
thin = Side(style='thin', color='FFBFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
F_TITLE = PatternFill('solid', fgColor='FF1F4E78')
F_CN = PatternFill('solid', fgColor='FF5B9BD5')
F_VN = PatternFill('solid', fgColor='FFD9EAF7')
F_NUM = PatternFill('solid', fgColor='FFEAF3F8')
F_IN = PatternFill('solid', fgColor='FFFFF9E6')    # input: light yellow
F_FX = PatternFill('solid', fgColor='FFF2F2F2')    # formula: light grey
F_HIN = PatternFill('solid', fgColor='FFFFE699')   # header of an input column
from openpyxl.worksheet.table import Table, TableStyleInfo, TableFormula
DATE = 'dd/mm/yyyy'; PCT = '0.0%'; INT = '#,##0'; TXT = '@'

wb = Workbook(); wb.remove(wb.active)

def sheet(name, title, cols, width_default=13):
    """cols: list of (cn, vn, width, kind['in'|'fx'], numfmt)"""
    ws = wb.create_sheet(name)
    ws['A1'] = title
    ws['A1'].font = Font(name=F0, size=14, bold=True, color='FFFFFFFF')
    ws['A1'].fill = F_TITLE
    ws['A1'].alignment = Alignment(vertical='center')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(len(cols), 8))
    ws.row_dimensions[1].height = 28; ws.row_dimensions[2].height = 28; ws.row_dimensions[3].height = 38
    for i, (cn, vn, w, kind, fmt) in enumerate(cols, 1):
        for r, txt, fill, font in ((2, cn, F_CN, Font(name=F0, size=10, bold=True, color='FFFFFFFF')),
                                   (3, vn, F_HIN if kind == 'in' else F_VN, Font(name=F0, size=10, bold=True, color='FF344767')),
                                   (4, str(i), F_NUM, Font(name=F0, size=9, bold=True, color='FF344767'))):
            c = ws.cell(r, i, txt); c.fill = fill; c.font = font; c.border = BORDER
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = 'A5'
    return ws

def body(ws, cols, first, last, formulas):
    """formulas: dict col_index -> function(row)->formula; other 'in' cols blank"""
    for r in range(first, last + 1):
        for i, (cn, vn, w, kind, fmt) in enumerate(cols, 1):
            c = ws.cell(r, i)
            if i in formulas: c.value = formulas[i](r)
            c.font = Font(name=F0, size=11)
            c.border = BORDER
            if fmt: c.number_format = fmt
            c.protection = Protection(locked=(kind != 'in'))
    t = Table(displayName="tbl" + ws.title.title().replace("_", ""), ref=f"A4:{L(len(cols))}{last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    t._initialise_columns()
    for i, tc in enumerate(t.tableColumns, 1):
        tc.name = str(i)
        if i in formulas:   # calculated column: Excel fills it into every new row typed under the table
            tc.calculatedColumnFormula = TableFormula(attr_text=formulas[i](first)[1:])
    ws.add_table(t)

def protect(ws):
    return  # no sheet protection: a protected sheet stops tables from growing
    p = ws.protection
    p.sheet = True
    p.formatColumns = False; p.formatRows = False; p.formatCells = False
    p.sort = False; p.autoFilter = False
    p.insertRows = True; p.deleteRows = True

def dname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

def dv_list(ws, formula, rng, allow_blank=True, msg=None):
    dv = DataValidation(type='list', formula1=formula, allow_blank=allow_blank)
    dv.error = msg or 'Chọn giá trị trong danh sách.'; dv.errorTitle = 'Sai dữ liệu'; dv.showErrorMessage = True
    ws.add_data_validation(dv); dv.add(rng)

# ---------------- CAU_HINH_CA ----------------
ca_cols = [('班次代碼','MÃ CA',12,'in',TXT),('班次名稱','TÊN CA',16,'in',None),('總工時','TỔNG GIỜ\n(tự tính)',12,'fx','0.0#'),
 ('第一時段開始','ĐỢT 1 BẮT ĐẦU',13,'in','hh:mm'),('第一時段結束','ĐỢT 1 KẾT THÚC',13,'in','hh:mm'),
 ('第二時段開始','ĐỢT 2 BẮT ĐẦU',13,'in','hh:mm'),('第二時段結束','ĐỢT 2 KẾT THÚC',13,'in','hh:mm'),
 ('第三時段開始','ĐỢT 3 BẮT ĐẦU',13,'in','hh:mm'),('第三時段結束','ĐỢT 3 KẾT THÚC',13,'in','hh:mm'),
 ('啟用狀態','TRẠNG THÁI',16,'in',None)]
ws = sheet('CAU_HINH_CA', '工作時間設定 / CẤU HÌNH CA', ca_cols)
body(ws, ca_cols, 5, 8, {3: lambda r: f'=IF(A{r}="","",ROUND(((E{r}-D{r})+(G{r}-F{r})+(I{r}-H{r}))*24,2))'})
T = dt.time
for k, row in enumerate([('8H','8 giờ',T(7,30),T(11,30),T(12,30),T(16,30),None,None),
                         ('9H','9 giờ',T(7,30),T(11,30),T(12,30),T(17,30),None,None),
                         ('10H','10 giờ',T(7,30),T(11,30),T(12,30),T(18,30),None,None),
                         ('11H30','11 giờ 30 phút',T(7,30),T(11,30),T(12,30),T(16,30),T(17,0),T(20,30))]):
    r = 5 + k
    ws.cell(r,1,row[0]); ws.cell(r,2,row[1])
    for j, v in enumerate(row[2:]):
        if v is not None: ws.cell(r, 4 + j, v)
    ws.cell(r,10,'Đang sử dụng')
dv_list(ws, '"Đang sử dụng,Ngừng sử dụng"', 'J5:J8')
protect(ws)
dname('DS_MA_CA', "OFFSET(CAU_HINH_CA!$A$5,0,0,MAX(1,COUNTA(CAU_HINH_CA!$A$5:$A$50)),1)")

# ---------------- DANH_SACH_CHUYEN ----------------
ch_cols = [('產線代碼','MÃ CHUYỀN',14,'in',TXT),('產線名稱','TÊN CHUYỀN',18,'in',None),('線長姓名','TÊN TRƯỞNG CHUYỀN',26,'in',None),('啟用狀態','TRẠNG THÁI',18,'in',None)]
ws = sheet('DANH_SACH_CHUYEN', '產線清單 / DANH SÁCH CHUYỀN', ch_cols)
body(ws, ch_cols, 5, LN, {})
for i in range(6):
    ws.cell(5+i,1,f'CH0{i+1}'); ws.cell(5+i,2,f'Chuyền {i+1}'); ws.cell(5+i,4,'Đang sử dụng')
dv_list(ws, '"Đang sử dụng,Ngừng sử dụng"', 'D5:D10')
dname('DS_CHUYEN', f"OFFSET(DANH_SACH_CHUYEN!$B$5,0,0,MAX(1,COUNTA(DANH_SACH_CHUYEN!$B$5:$B${LR})),1)")

# ---------------- DANH_SACH_SAN_PHAM ----------------
sp_cols = [('产品代码','MÃ SẢN PHẨM',20,'in',TXT),('产品名称','TÊN SẢN PHẨM',28,'in',None),('備註','GHI CHÚ',28,'in',None),('啟用狀態','TRẠNG THÁI',18,'in',None)]
ws = sheet('DANH_SACH_SAN_PHAM', '產品清單 / DANH SÁCH SẢN PHẨM', sp_cols)
body(ws, sp_cols, 5, 9, {})
for i, (code, note) in enumerate([('ĐAI LƯNG',''),('BAO TAY XANH',''),('883',''),('567','Test'),('MÃ 9H','Mã thử từ V17')]):
    ws.cell(5+i,1,code); ws.cell(5+i,3,note or None); ws.cell(5+i,4,'Đang sử dụng')
dv_list(ws, '"Đang sử dụng,Ngừng sử dụng"', 'D5:D9')
dname('DS_MA_SAN_PHAM', "OFFSET(DANH_SACH_SAN_PHAM!$A$5,0,0,MAX(1,COUNTA(DANH_SACH_SAN_PHAM!$A$5:$A$1000)),1)")

# ---------------- DANH_SACH_LY_DO ----------------
ld_cols = [('原因','LÝ DO KHÔNG ĐẠT',28,'in',None),('備註','GHI CHÚ',36,'in',None)]
ws = sheet('DANH_SACH_LY_DO', '未達原因清單 / DANH SÁCH LÝ DO KHÔNG ĐẠT', ld_cols)
REASONS = ['Thiếu vật tư','Hư máy','Thiếu người','Chờ hàng / chờ khuôn','Sửa hàng lỗi','Khác']
body(ws, ld_cols, 5, 4 + len(REASONS), {})
for i, t in enumerate(REASONS): ws.cell(5 + i, 1, t)
dname('DS_LY_DO', "OFFSET(DANH_SACH_LY_DO!$A$5,0,0,MAX(1,COUNTA(DANH_SACH_LY_DO!$A$5:$A$1000)),1)")

# ---------------- DANH_SACH_LOAI_LOI (V20) ----------------
ll2_cols = [('不良類型','LOẠI LỖI',28,'in',None),('備註','GHI CHÚ',36,'in',None)]
ws = sheet('DANH_SACH_LOAI_LOI', '不良類型清單 / DANH SÁCH LOẠI LỖI', ll2_cols)
DEFECTS = ['Rách đường may','Bung chỉ','Lem màu','Sai kích thước','Khóa / phụ kiện lệch','Khác']
body(ws, ll2_cols, 5, 4 + len(DEFECTS), {})
for i, t in enumerate(DEFECTS): ws.cell(5 + i, 1, t)
dname('DS_LOAI_LOI', "OFFSET(DANH_SACH_LOAI_LOI!$A$5,0,0,MAX(1,COUNTA(DANH_SACH_LOAI_LOI!$A$5:$A$1000)),1)")

# ---------------- LICH_LAM_VIEC ----------------
ll_cols = [('日期','NGÀY',12,'in',DATE),('類型','LOẠI',14,'in',None),('星期','THỨ',8,'fx',None),('備註','GHI CHÚ',30,'in',None)]
ws = sheet('LICH_LAM_VIEC', '工作日曆 / LỊCH LÀM VIỆC (Chủ nhật mặc định nghỉ)', ll_cols)
body(ws, ll_cols, 5, 5, {3: lambda r: f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r}),"CN","T2","T3","T4","T5","T6","T7"))'})
ws.cell(5, 1, dt.datetime(2027, 1, 1)); ws.cell(5, 2, 'Nghỉ'); ws.cell(5, 4, 'Tết Dương lịch (ví dụ)')
dv_list(ws, '"Nghỉ,Làm bù"', 'B5:B5', msg='Chọn Nghỉ (ngày thường được nghỉ) hoặc Làm bù (Chủ nhật đi làm).')
LLA = 'LICH_LAM_VIEC!$A$5:$A$1000'; LLB = 'LICH_LAM_VIEC!$B$5:$B$1000'
def WD(a, b):
    """working days a..b: Mon–Sat, minus 'Nghỉ' weekdays, plus 'Làm bù' Sundays"""
    return (f'MAX(0,NETWORKDAYS.INTL({a},{b},11)'
            f'-SUMPRODUCT(({LLA}>={a})*({LLA}<={b})*({LLB}="Nghỉ")*(WEEKDAY({LLA})<>1))'
            f'+SUMPRODUCT(({LLA}>={a})*({LLA}<={b})*({LLB}="Làm bù")*(WEEKDAY({LLA})=1)))')

# ---------------- NHAP_LIEU ----------------
R = lambda c: f"NHAP_LIEU!${c}$5:${c}${N}"
HLR = lambda c: f"HANG_LOI!${c}$5:${c}${HN}"   # defect log (V20)
ANYH = '((' + '+'.join(f'(NHAP_LIEU!${c}$5:${c}${N}<>"")' for c in 'HIJKLMNOPQRS') + ')>0)'   # row has any hour typed
NHRS = '(' + '+'.join(f'(NHAP_LIEU!${c}$5:${c}${N}<>"")' for c in 'HIJKLMNOPQRS') + ')'         # hours typed per row
CAT = f'ROUND(SUMIF(CAU_HINH_CA!$A$5:$A$50,NHAP_LIEU!$D$5:$D${N},CAU_HINH_CA!$C$5:$C$50)*NHAP_LIEU!$E$5:$E${N},0)'  # daily target per row
nl_cols = [('日期','NGÀY',12,'in',DATE),('產線','CHUYỀN',12,'in',None),('产品代码','MÃ SẢN PHẨM',18,'in',TXT),
 ('工作時間','MÃ CA',10,'in',TXT),('每小时目标产量PCS','MỤC TIÊU MỖI GIỜ',12,'in',INT),
 ('總工時','GIỜ CA',9,'fx','0.0#'),('每日目标产量PCS','MỤC TIÊU TRONG NGÀY',13,'fx',INT)]
for h in range(1, 13):
    nl_cols.append((f'第{h}小時', f'GIỜ {h}', 8, 'in', INT))
nl_cols += [('每日实际产量PCS','THỰC TẾ TRONG NGÀY',13,'fx',INT),('当日达成率%','TỶ LỆ ĐẠT TRONG NGÀY',11,'fx',PCT),
 ('差異PCS','CHÊNH LỆCH',11,'fx','#,##0;[Red]-#,##0'),('剩餘PCS','CÒN THIẾU',11,'fx',INT),
 ('已輸入小時數','SỐ GIỜ ĐÃ NHẬP',9,'fx','0'),('累計目標PCS','MỤC TIÊU ĐẾN GIỜ ĐÃ NHẬP',13,'fx',INT),
 ('進度達成率%','TIẾN ĐỘ THEO GIỜ',11,'fx',PCT),('前一工作日','NGÀY LÀM TRƯỚC',12,'fx','dd/mm/yyyy;;'),
 ('前一日欠數PCS','THIẾU HÔM TRƯỚC',12,'fx',INT),('狀態','TRẠNG THÁI',13,'in',None),('備註','GHI CHÚ',22,'in',None),
 ('檢查','KIỂM TRA',22,'fx',None),('索引','KHÓA (máy dùng)',16,'fx',None),
 ('工人數','SỐ CÔNG NHÂN',10,'in','0'),('不良數PCS','SỐ LỖI\n(tự cộng từ HANG_LOI)',10,'fx',INT),('未達原因','LÝ DO KHÔNG ĐẠT',22,'in',None),('停機分鐘','PHÚT DỪNG MÁY',10,'in','0'),
 ('人均每小時產量','SP / NGƯỜI / GIỜ',11,'fx','0.0'),('不良率%','TỶ LỆ LỖI',10,'fx','0.00%')]
ws = sheet('NHAP_LIEU', '生產計劃與每小時產量 / NHẬP LIỆU HẰNG NGÀY (KẾ HOẠCH + SẢN LƯỢNG THEO GIỜ)', nl_cols)
fx = {
 6: lambda r: f'=IF(D{r}="","",IFERROR(INDEX(CAU_HINH_CA!$C$5:$C$50,MATCH(D{r},CAU_HINH_CA!$A$5:$A$50,0)),""))',
 7: lambda r: f'=IF(OR(F{r}="",E{r}=""),"",ROUND(F{r}*E{r},0))',
 20: lambda r: f'=IF(COUNT(H{r}:S{r})=0,"",SUM(H{r}:S{r}))',
 21: lambda r: f'=IF(OR(T{r}="",N(G{r})=0),"",T{r}/G{r})',
 22: lambda r: f'=IF(OR(T{r}="",G{r}=""),"",T{r}-G{r})',
 23: lambda r: f'=IF(OR(T{r}="",G{r}=""),"",MAX(0,G{r}-T{r}))',
 24: lambda r: f'=IF(A{r}="","",COUNT(H{r}:S{r}))',
 25: lambda r: f'=IF(OR(X{r}="",E{r}="",G{r}=""),"",MIN(G{r},X{r}*E{r}))',
 26: lambda r: f'=IF(OR(T{r}="",N(Y{r})=0),"",T{r}/Y{r})',
 27: lambda r: (lambda p: f'=IF(OR(A{r}="",B{r}=""),"",IF({p}=0,"",{p}))')(f'SUMPRODUCT(MAX(($B$5:$B${N}=B{r})*($A$5:$A${N}<A{r})*$A$5:$A${N}))'),
 28: lambda r: f'=IF(N(AA{r})=0,"",SUMIFS($W$5:$W${N},$B$5:$B${N},B{r},$A$5:$A${N},AA{r}))',
 31: lambda r: (f'=IF(OR(A{r}="",B{r}=""),"",IF(COUNTIFS($A$5:$A${N},A{r},$B$5:$B${N},B{r})>1,"TRÙNG NGÀY + CHUYỀN",'
                f'IF(G{r}="","THIẾU MÃ CA / MỤC TIÊU GIỜ",'
                f'IF(SUMPRODUCT((COLUMN(H{r}:S{r})-COLUMN(H{r})+1>ROUNDUP(F{r},0))*(H{r}:S{r}<>""))>0,"NHẬP QUÁ GIỜ CA",'
                f'IF(AND(N(X{r})>=ROUNDUP(N(F{r}),0),N(F{r})>0,N(T{r})<N(G{r}),AI{r}=""),"CHƯA GHI LÝ DO KHÔNG ĐẠT",'
                f'IF(COUNTIFS(MUC_TIEU_THANG!$B$5:$B${MN},C{r},MUC_TIEU_THANG!$A$5:$A${MN},">="&DATE(YEAR(A{r}),MONTH(A{r}),1),MUC_TIEU_THANG!$A$5:$A${MN},"<"&DATE(YEAR(A{r}),MONTH(A{r})+1,1))=0,"CHƯA CÓ MỤC TIÊU THÁNG","OK"))))))'),
 32: lambda r: f'=IF(OR(A{r}="",B{r}=""),"",B{r}&"|"&A{r})',
 34: lambda r: f'=IF(OR(A{r}="",B{r}="",COUNTIFS({HLR("A")},A{r},{HLR("C")},B{r})=0),"",SUMIFS({HLR("F")},{HLR("A")},A{r},{HLR("C")},B{r}))',
 37: lambda r: f'=IF(OR(T{r}="",N(AG{r})=0,N(X{r})=0),"",T{r}/(AG{r}*X{r}))',
 38: lambda r: f'=IF(OR(AH{r}="",N(T{r})=0),"",AH{r}/T{r})',
}
body(ws, nl_cols, 5, 9, fx)
ws.freeze_panes = 'D5'
for k, (line, prod, ca, tph) in enumerate([('Chuyền 1','ĐAI LƯNG','8H',100),('Chuyền 2','BAO TAY XANH','10H',100),
        ('Chuyền 3','883','11H30',185),('Chuyền 4','MÃ 9H','9H',100),('Chuyền 6','567','10H',40)]):
    r = 5 + k
    ws.cell(r,1,dt.datetime(2026,9,30)); ws.cell(r,2,line); ws.cell(r,3,prod); ws.cell(r,4,ca); ws.cell(r,5,tph); ws.cell(r,29,'Đang chạy')
if SAMPLE:
    for r, vals in [(5,[100,95,110,90,100,98,105,92]),(6,[230,240,236,228,232,240,235,240,236,233]),(7,[180,170,185,190,186]),
                    (8,[100,96,104]),(9,[40,42,38,40,41,39])]:
        for i, v in enumerate(vals): ws.cell(r, 8 + i, v)
dv_list(ws, 'DS_CHUYEN', f'B5:B{NL}')
dv_list(ws, 'DS_MA_SAN_PHAM', f'C5:C{NL}', msg='Mã sản phẩm phải có trong sheet DANH_SACH_SAN_PHAM.')
dv_list(ws, 'DS_MA_CA', f'D5:D{NL}')
dv_list(ws, '"Chưa chạy,Đang chạy,Hoàn thành,Tạm dừng"', f'AC5:AC{NL}')
dvd = DataValidation(type='date', operator='greaterThan', formula1='DATE(2020,1,1)', allow_blank=True)
dvd.error='Nhập ngày dạng dd/mm/yyyy.'; dvd.showErrorMessage=True; ws.add_data_validation(dvd); dvd.add(f'A5:A{NL}')
dvn = DataValidation(type='whole', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
dvn.error='Nhập số nguyên ≥ 0.'; dvn.showErrorMessage=True; ws.add_data_validation(dvn); dvn.add(f'E5:E{NL}'); dvn.add(f'H5:S{NL}'); dvn.add(f'AG5:AG{NL}'); dvn.add(f'AJ5:AJ{NL}')
dv_list(ws, 'DS_LY_DO', f'AI5:AI{NL}', msg='Chọn lý do trong sheet DANH_SACH_LY_DO.')
red = PatternFill('solid', fgColor='FFF8CBAD'); green = PatternFill('solid', fgColor='FFC6EFCE')
CA_HRS = 'IFERROR(INDEX(CAU_HINH_CA!$C$5:$C$50,MATCH($D5,CAU_HINH_CA!$A$5:$A$50,0)),99)'
ws.conditional_formatting.add(f'A5:AL{NL}', FormulaRule(formula=[f'AND($A5<>"",$B5<>"",COUNTIFS($A$5:$A${N},$A5,$B$5:$B${N},$B5)>1)'], fill=red))
ws.conditional_formatting.add(f'H5:S{NL}', FormulaRule(formula=[f'AND(H5<>"",COLUMN(H5)-COLUMN($H5)+1>ROUNDUP({CA_HRS},0))'], fill=PatternFill('solid', fgColor='FFFF7C80'), font=Font(color='FF9C0006', bold=True)))
ws.conditional_formatting.add(f'AE5:AE{NL}', FormulaRule(formula=[f'AND(AE5<>"",AE5<>"OK")'], fill=red))
for col in ('U', 'Z'):
    ws.conditional_formatting.add(f'{col}5:{col}{NL}', FormulaRule(formula=[f'AND(ISNUMBER({col}5),{col}5<1)'], font=Font(color='FFC00000', bold=True)))
    ws.conditional_formatting.add(f'{col}5:{col}{NL}', FormulaRule(formula=[f'AND(ISNUMBER({col}5),{col}5>=1)'], font=Font(color='FF007A33', bold=True)))
ws.conditional_formatting.add(f'AB5:AB{NL}', FormulaRule(formula=['AND(ISNUMBER(AB5),AB5>0)'], font=Font(color='FFC00000', bold=True)))
ws.column_dimensions['AF'].hidden = True
protect(ws)

# ---------------- HANG_LOI (V20): one row per defect batch, with a photo ----------------
hl_cols = [('日期','NGÀY',12,'in',DATE),('時間','GIỜ',8,'in','hh:mm'),('產線','CHUYỀN',12,'in',None),
 ('产品代码','MÃ SẢN PHẨM\n(tự lấy)',18,'fx',TXT),('不良類型','LOẠI LỖI',22,'in',None),('數量PCS','SỐ LƯỢNG',10,'in',INT),
 ('照片檔名','TÊN FILE ẢNH',26,'in',None),('備註','GHI CHÚ',24,'in',None),
 ('電視顯示','HIỆN TRÊN TV',10,'fx',None),('檢查','KIỂM TRA',26,'fx',None)]
ws = sheet('HANG_LOI', '不良品記錄 / HÀNG LỖI (kèm ảnh, hiện trên TV)', hl_cols)
HL_SEED = [(T(10,15),'Chuyền 2','Rách đường may',1,'vi_du_rach.jpg'),
           (T(9,30),'Chuyền 1','Khóa / phụ kiện lệch',2,None)]
if SAMPLE:
    HL_SEED = [(T(8,50),'Chuyền 3','Bung chỉ',4,'bung_chi_c3.png'),
               (T(9,30),'Chuyền 1','Khóa / phụ kiện lệch',12,'khoa_lech_c1.png'),
               (T(10,15),'Chuyền 2','Rách đường may',40,'rach_c2.png'),
               (T(11,5),'Chuyền 4','Sai kích thước',21,'sai_size_c4.png'),
               (T(13,40),'Chuyền 2','Lem màu',45,'lem_mau_c2.png'),
               (T(14,20),'Chuyền 6','Khác',3,None)]
HLN = 4 + len(HL_SEED)
body(ws, hl_cols, 5, HLN, {
 4: lambda r: f'=IF(OR(A{r}="",C{r}=""),"",IFERROR(INDEX({R("C")},MATCH(1,INDEX(({R("A")}=A{r})*({R("B")}=C{r}),0),0))&"",""))',
 9: lambda r: f'=IF(OR(A{r}="",C{r}=""),"",IF(A{r}=HIEN_THI!$M$1,"CÓ",""))',
 10: lambda r: (f'=IF(AND(A{r}="",C{r}=""),"",IF(OR(A{r}="",C{r}=""),"THIẾU NGÀY / CHUYỀN",IF(D{r}="","CHƯA CÓ DÒNG NHẬP LIỆU NGÀY NÀY",'
                f'IF(N(F{r})<=0,"THIẾU SỐ LƯỢNG",IF(E{r}="","CHƯA CHỌN LOẠI LỖI",IF(G{r}="","OK (chưa có ảnh)","OK"))))))'),
})
for k, (h, line, typ, qty, img) in enumerate(HL_SEED):
    r = 5 + k
    ws.cell(r,1,dt.datetime(2026,9,30)); ws.cell(r,2,h); ws.cell(r,3,line); ws.cell(r,5,typ)
    if qty: ws.cell(r,6,qty)
    if img: ws.cell(r,7,img)
dv_list(ws, 'DS_CHUYEN', f'C5:C{HLN}')
dv_list(ws, 'DS_LOAI_LOI', f'E5:E{HLN}', msg='Chọn loại lỗi trong sheet DANH_SACH_LOAI_LOI.')
dvd2 = DataValidation(type='date', operator='greaterThan', formula1='DATE(2020,1,1)', allow_blank=True)
dvd2.error='Nhập ngày dạng dd/mm/yyyy.'; dvd2.showErrorMessage=True; ws.add_data_validation(dvd2); dvd2.add(f'A5:A{HLN}')
dvq = DataValidation(type='whole', operator='greaterThan', formula1='0', allow_blank=True)
dvq.error='Nhập số nguyên > 0.'; dvq.showErrorMessage=True; ws.add_data_validation(dvq); dvq.add(f'F5:F{HLN}')
ws.conditional_formatting.add(f'J5:J{HLN}', FormulaRule(formula=[f'AND(J5<>"",LEFT(J5,2)<>"OK")'], fill=PatternFill('solid', fgColor='FFF8CBAD')))
ws.conditional_formatting.add(f'I5:I{HLN}', FormulaRule(formula=['I5="CÓ"'], font=Font(color='FF007A33', bold=True)))
ws.freeze_panes = 'D5'

# ---------------- MUC_TIEU_THANG ----------------
mt_cols = [('月份','THÁNG',12,'in','mm/yyyy'),('产品代码','MÃ SẢN PHẨM',18,'in',TXT),('当月总目标产量PCS','TỔNG SẢN LƯỢNG MỤC TIÊU TRONG THÁNG',16,'in',INT),
 ('備註','GHI CHÚ',20,'in',None),('当月累计产能PCS','LŨY KẾ SẢN LƯỢNG TRONG THÁNG',16,'fx',INT),
 ('当月累计差异PCS','LŨY KẾ CHÊNH LỆCH TRONG THÁNG',16,'fx','#,##0;[Red]-#,##0'),('当月总达成率','TỔNG ĐẠT % TRONG THÁNG',12,'fx',PCT),
 ('剩餘PCS','CÒN THIẾU TRONG THÁNG',14,'fx',INT),('累計每日目標PCS','LŨY KẾ MỤC TIÊU NGÀY (các ngày đã nhập)',16,'fx',INT),
 ('與每日目標差異PCS','CHÊNH LỆCH SO VỚI MỤC TIÊU NGÀY',16,'fx','#,##0;[Red]-#,##0'),
 ('上月欠數PCS','THIẾU THÁNG TRƯỚC',14,'fx',INT),('當月工作天數','SỐ NGÀY LÀM VIỆC TRONG THÁNG',12,'fx','0'),('平均每日需產量PCS','BÌNH QUÂN CẦN / NGÀY',14,'fx',INT)]
ws = sheet('MUC_TIEU_THANG', '當月目標與累計 / MỤC TIÊU & LŨY KẾ THÁNG', mt_cols)
M = lambda r: (f'{R("C")},B{r},{R("A")},">="&DATE(YEAR(A{r}),MONTH(A{r}),1),{R("A")},"<"&DATE(YEAR(A{r}),MONTH(A{r})+1,1)')
body(ws, mt_cols, 5, 8, {
 5: lambda r: f'=IF(OR(A{r}="",B{r}=""),"",SUMPRODUCT(({R("C")}=B{r})*({R("A")}>=DATE(YEAR(A{r}),MONTH(A{r}),1))*({R("A")}<DATE(YEAR(A{r}),MONTH(A{r})+1,1))*NHAP_LIEU!$H$5:$S${N}))',
 6: lambda r: f'=IF(OR(E{r}="",C{r}=""),"",E{r}-C{r})',
 7: lambda r: f'=IF(OR(E{r}="",N(C{r})=0),"",E{r}/C{r})',
 8: lambda r: f'=IF(OR(E{r}="",C{r}=""),"",MAX(0,C{r}-E{r}))',
 9: lambda r: f'=IF(OR(A{r}="",B{r}=""),"",SUMPRODUCT(({R("C")}=B{r})*({R("A")}>=DATE(YEAR(A{r}),MONTH(A{r}),1))*({R("A")}<DATE(YEAR(A{r}),MONTH(A{r})+1,1))*(((NHAP_LIEU!$H$5:$H${N}<>"")+(NHAP_LIEU!$I$5:$I${N}<>"")+(NHAP_LIEU!$J$5:$J${N}<>"")+(NHAP_LIEU!$K$5:$K${N}<>"")+(NHAP_LIEU!$L$5:$L${N}<>"")+(NHAP_LIEU!$M$5:$M${N}<>"")+(NHAP_LIEU!$N$5:$N${N}<>"")+(NHAP_LIEU!$O$5:$O${N}<>"")+(NHAP_LIEU!$P$5:$P${N}<>"")+(NHAP_LIEU!$Q$5:$Q${N}<>"")+(NHAP_LIEU!$R$5:$R${N}<>"")+(NHAP_LIEU!$S$5:$S${N}<>""))>0)*ROUND(SUMIF(CAU_HINH_CA!$A$5:$A$50,{R("D")},CAU_HINH_CA!$C$5:$C$50)*{R("E")},0)))',
 10: lambda r: f'=IF(I{r}="","",E{r}-I{r})',
 11: lambda r: f'=IF(OR(A{r}="",B{r}=""),"",SUMIFS($H$5:$H${MN},$B$5:$B${MN},B{r},$A$5:$A${MN},">="&DATE(YEAR(A{r}),MONTH(A{r})-1,1),$A$5:$A${MN},"<"&DATE(YEAR(A{r}),MONTH(A{r}),1)))',
 12: lambda r: '=IF(A{0}="","",{1})'.format(r, WD(f'DATE(YEAR(A{r}),MONTH(A{r}),1)', f'EOMONTH(A{r},0)')),
 13: lambda r: f'=IF(OR(N(L{r})=0,C{r}=""),"",ROUNDUP(C{r}/L{r},0))',
})
for k, (prod, tgt, note) in enumerate([('ĐAI LƯNG',40000,None),('BAO TAY XANH',40000,None),('883',50000,None),('567',100000,'Test')]):
    r = 5 + k; ws.cell(r,1,dt.datetime(2026,9,1)); ws.cell(r,2,prod); ws.cell(r,3,tgt); ws.cell(r,4,note)
dv_list(ws, 'DS_MA_SAN_PHAM', 'B5:B8', msg='Mã sản phẩm phải có trong sheet DANH_SACH_SAN_PHAM.')
ws.conditional_formatting.add('G5:G8', FormulaRule(formula=['AND(ISNUMBER(G5),G5<1)'], font=Font(color='FFC00000', bold=True)))
protect(ws)

# ---------------- HIEN_THI ----------------
ht_cols = [('日期','NGÀY',12,'fx',DATE),('產線','CHUYỀN',12,'fx',None),('产品代码','MÃ SẢN PHẨM',16,'fx',TXT),('工作時間','MÃ CA',9,'fx',None),
 ('總工時','GIỜ CA',8,'fx','0.0#'),('每小时目标产量PCS','MỤC TIÊU MỖI GIỜ',11,'fx',INT),('每日目标产量PCS','MỤC TIÊU TRONG NGÀY',12,'fx',INT),
 ('每日实际产量PCS','THỰC TẾ TRONG NGÀY',12,'fx',INT),('当日达成率%','TỶ LỆ ĐẠT TRONG NGÀY',11,'fx',PCT),
 ('差異PCS','CHÊNH LỆCH',11,'fx','#,##0;[Red]-#,##0'),('剩餘PCS','CÒN THIẾU',11,'fx',INT),('已輸入小時數','SỐ GIỜ ĐÃ NHẬP',9,'fx','0'),
 ('累計目標PCS','MỤC TIÊU ĐẾN GIỜ ĐÃ NHẬP',12,'fx',INT),('進度達成率%','TIẾN ĐỘ THEO GIỜ',11,'fx',PCT),
 ('前一工作日','NGÀY LÀM TRƯỚC',12,'fx','dd/mm/yyyy;;'),('前一日欠數PCS','THIẾU HÔM TRƯỚC',12,'fx',INT),
 ('当月总目标产量PCS','MỤC TIÊU THÁNG (SẢN PHẨM)',13,'fx',INT),('当月累计产能PCS','LŨY KẾ THÁNG (SẢN PHẨM, đến ngày này)',14,'fx',INT),
 ('当月总达成率','TỶ LỆ ĐẠT THÁNG',11,'fx',PCT),('剩餘PCS','CÒN THIẾU THÁNG',12,'fx',INT),('本線當月累計PCS','LŨY KẾ THÁNG CỦA CHUYỀN',13,'fx',INT)]
for h in range(1, 13):
    ht_cols.append((f'第{h}小時', f'GIỜ {h}', 8, 'fx', INT))
ht_cols += [('狀態','TRẠNG THÁI',12,'fx',None),('備註','GHI CHÚ',20,'fx',None),('索引','DÒNG NHẬP LIỆU (máy dùng)',10,'fx','0'),('前一日索引','DÒNG NGÀY TRƯỚC (máy dùng)',10,'fx','0'),
 ('上月欠數PCS','THIẾU THÁNG TRƯỚC (SẢN PHẨM)',13,'fx',INT),('剩餘工作天數','NGÀY LÀM VIỆC CÒN LẠI (tính cả ngày này)',11,'fx','0'),
 ('每日需產量PCS','CẦN LÀM MỖI NGÀY ĐỂ KỊP THÁNG (SẢN PHẨM)',14,'fx',INT),('工人數','SỐ CÔNG NHÂN',10,'fx','0'),
 ('人均每小時產量','SP / NGƯỜI / GIỜ',11,'fx','0.0'),('不良數PCS','SỐ LỖI',10,'fx',INT),('不良率%','TỶ LỆ LỖI',10,'fx','0.00%'),
 ('未達原因','LÝ DO KHÔNG ĐẠT',20,'fx',None),('停機分鐘','PHÚT DỪNG MÁY',10,'fx','0')]
ws = sheet('HIEN_THI', '電視顯示 / BẢNG HIỂN THỊ TV', ht_cols)
ws.unmerge_cells('A1:H1'); ws.merge_cells('A1:E1')
lab = Font(name=F0, size=11, bold=True, color='FF1F4E78')
ws['G1'] = 'CHỌN NGÀY (để trống = ngày mới nhất):'; ws['G1'].font = lab; ws['G1'].alignment = Alignment(horizontal='right'); ws.merge_cells('G1:I1')
ws['J1'].fill = F_IN; ws['J1'].border = BORDER; ws['J1'].number_format = DATE; ws['J1'].protection = Protection(locked=False)
ws['K1'] = 'ĐANG HIỂN THỊ:'; ws['K1'].font = lab; ws['K1'].alignment = Alignment(horizontal='right'); ws.merge_cells('K1:L1')
ws['M1'] = f'=IF(J1<>"",J1,MAX({R("A")}))'
ws['M1'].number_format = DATE; ws['M1'].font = Font(name=F0, size=12, bold=True, color='FFC00000'); ws['M1'].fill = F_FX; ws['M1'].border = BORDER
X = 'AJ'   # matched NHAP_LIEU row (found from the typed date + line only)
def g(col, r):  # typed value from NHAP_LIEU at the matched row
    return f'=IF(${X}{r}="","",IF(INDEX({R(col)},${X}{r})="","",INDEX({R(col)},${X}{r})))'
MS = lambda r: f'DATE(YEAR(A{r}),MONTH(A{r}),1)'
HS = f'NHAP_LIEU!$H$5:$S${N}'          # hourly input block
CA_H = lambda code: f'INDEX(CAU_HINH_CA!$C$5:$C$50,MATCH({code},CAU_HINH_CA!$A$5:$A$50,0))'
# HIEN_THI reads only the typed columns of NHAP_LIEU (A–E, H–S), never its formula columns,
# so a new day shows even if a formula column was not filled down on the new row.
hx = {
 1: lambda r: f'=IF(B{r}="","",$M$1)',
 2: lambda r: f'=IF(DANH_SACH_CHUYEN!B{r}="","",DANH_SACH_CHUYEN!B{r})',
 3: lambda r: g('C', r), 4: lambda r: g('D', r),
 5: lambda r: f'=IF(D{r}="","",IFERROR({CA_H(f"D{r}")},""))',
 6: lambda r: g('E', r),
 7: lambda r: f'=IF(OR(E{r}="",F{r}=""),"",ROUND(E{r}*F{r},0))',
 8: lambda r: f'=IF(COUNT(V{r}:AG{r})=0,"",SUM(V{r}:AG{r}))',
 9: lambda r: f'=IF(OR(H{r}="",N(G{r})=0),"",H{r}/G{r})',
 10: lambda r: f'=IF(OR(H{r}="",G{r}=""),"",H{r}-G{r})',
 11: lambda r: f'=IF(OR(H{r}="",G{r}=""),"",MAX(0,G{r}-H{r}))',
 12: lambda r: f'=IF({X}{r}="","",COUNT(V{r}:AG{r}))',
 13: lambda r: f'=IF(OR(L{r}="",F{r}="",G{r}=""),"",MIN(G{r},L{r}*F{r}))',
 14: lambda r: f'=IF(OR(H{r}="",N(M{r})=0),"",H{r}/M{r})',
 15: lambda r: (lambda p: f'=IF(OR(B{r}="",A{r}=""),"",IF({p}=0,"",{p}))')(f'SUMPRODUCT(MAX(({R("B")}=B{r})*({R("A")}<A{r})*{R("A")}))'),
 16: lambda r: (f'=IF(AK{r}="","",IF(COUNT(INDEX({HS},AK{r},0))=0,0,MAX(0,IFERROR(ROUND('
               f'{CA_H(f"INDEX({R(chr(68))},AK{r})")}*INDEX({R("E")},AK{r}),0),0)-SUM(INDEX({HS},AK{r},0)))))'),
 17: lambda r: f'=IF(C{r}="","",SUMIFS(MUC_TIEU_THANG!$C$5:$C${MN},MUC_TIEU_THANG!$B$5:$B${MN},C{r},MUC_TIEU_THANG!$A$5:$A${MN},">="&{MS(r)},MUC_TIEU_THANG!$A$5:$A${MN},"<"&DATE(YEAR(A{r}),MONTH(A{r})+1,1)))',
 18: lambda r: f'=IF(C{r}="","",SUMPRODUCT(({R("C")}=C{r})*({R("A")}>={MS(r)})*({R("A")}<=A{r})*{HS}))',
 19: lambda r: f'=IF(OR(R{r}="",N(Q{r})=0),"",R{r}/Q{r})',
 20: lambda r: f'=IF(OR(R{r}="",N(Q{r})=0),"",MAX(0,Q{r}-R{r}))',
 21: lambda r: f'=IF(B{r}="","",SUMPRODUCT(({R("B")}=B{r})*({R("A")}>={MS(r)})*({R("A")}<=A{r})*{HS}))',
 34: lambda r: g('AC', r), 35: lambda r: g('AD', r),
 36: lambda r: f'=IF(OR(B{r}="",A{r}=""),"",IFERROR(MATCH(1,INDEX(({R("A")}=A{r})*({R("B")}=B{r}),0),0),""))',
 37: lambda r: f'=IF(O{r}="","",IFERROR(MATCH(1,INDEX(({R("A")}=O{r})*({R("B")}=B{r}),0),0),""))',
 38: lambda r: f'=IF(C{r}="","",SUMIFS(MUC_TIEU_THANG!$K$5:$K${MN},MUC_TIEU_THANG!$B$5:$B${MN},C{r},MUC_TIEU_THANG!$A$5:$A${MN},">="&{MS(r)},MUC_TIEU_THANG!$A$5:$A${MN},"<"&DATE(YEAR(A{r}),MONTH(A{r})+1,1)))',
 39: lambda r: '=IF(A{0}="","",{1})'.format(r, WD(f'A{r}', f'EOMONTH(A{r},0)')),
 40: lambda r: f'=IF(OR(C{r}="",N(Q{r})=0,N(AM{r})=0),"",ROUNDUP(MAX(0,Q{r}-SUMPRODUCT(({R("C")}=C{r})*({R("A")}>={MS(r)})*({R("A")}<A{r})*{HS}))/AM{r},0))',
 41: lambda r: g('AG', r),
 42: lambda r: f'=IF(OR(H{r}="",N(AO{r})=0,N(L{r})=0),"",H{r}/(AO{r}*L{r}))',
 43: lambda r: f'=IF(OR(B{r}="",A{r}="",COUNTIFS({HLR("A")},A{r},{HLR("C")},B{r})=0),"",SUMIFS({HLR("F")},{HLR("A")},A{r},{HLR("C")},B{r}))',
 44: lambda r: f'=IF(OR(AQ{r}="",N(H{r})=0),"",AQ{r}/H{r})',
 45: lambda r: g('AI', r), 46: lambda r: g('AJ', r),
}
src = ['H','I','J','K','L','M','N','O','P','Q','R','S']
for h in range(12):
    hx[22 + h] = (lambda c: (lambda r: g(c, r)))(src[h])
body(ws, ht_cols, 5, LN, hx)
# total row
TR = LN + 2
ws.cell(TR, 2, 'TỔNG CỘNG')
tot = {7:'G',8:'H',10:'J',11:'K',13:'M',16:'P',21:'U',41:'AO',43:'AQ',46:'AT'}
for i in range(1, len(ht_cols) + 1):
    c = ws.cell(TR, i); c.font = Font(name=F0, size=11, bold=True); c.fill = F_VN; c.border = BORDER
    c.number_format = ht_cols[i-1][4] or 'General'
for i, col in tot.items():
    ws.cell(TR, i, f'=IF(COUNT({col}5:{col}{LN})=0,"",SUM({col}5:{col}{LN}))')
ws.cell(TR, 1, '=$M$1'); ws.cell(TR, 1).number_format = DATE
ws.cell(TR, 9, f'=IF(OR(H{TR}="",N(G{TR})=0),"",H{TR}/G{TR})')
ws.cell(TR, 14, f'=IF(OR(H{TR}="",N(M{TR})=0),"",H{TR}/M{TR})')
ws.cell(TR, 42, f'=IF(SUMPRODUCT(AO5:AO{LN},L5:L{LN})=0,"",H{TR}/SUMPRODUCT(AO5:AO{LN},L5:L{LN}))')
ws.cell(TR, 44, f'=IF(OR(AQ{TR}="",N(H{TR})=0),"",AQ{TR}/H{TR})')
for col in ('I', 'N', 'S'):
    ws.conditional_formatting.add(f'{col}5:{col}{TR}', FormulaRule(formula=[f'AND(ISNUMBER({col}5),{col}5<1)'], font=Font(color='FFC00000', bold=True)))
    ws.conditional_formatting.add(f'{col}5:{col}{TR}', FormulaRule(formula=[f'AND(ISNUMBER({col}5),{col}5>=1)'], font=Font(color='FF007A33', bold=True)))
protect(ws)

# ---------------- BAO_CAO_THANG (fixed-size monthly report, not a table) ----------------
from openpyxl.chart import BarChart, LineChart, Reference
ws = wb.create_sheet('BAO_CAO_THANG')
HDR = Font(name=F0, size=10, bold=True, color='FF344767'); BOLD = Font(name=F0, size=11, bold=True)
SEC = Font(name=F0, size=12, bold=True, color='FF1F4E78')
def cellf(r, c, v=None, fmt=None, font=None, fill=None, border=True, align=None):
    x = ws.cell(r, c)
    if v is not None: x.value = v
    if fmt: x.number_format = fmt
    x.font = font or Font(name=F0, size=11)
    if fill: x.fill = fill
    if border: x.border = BORDER
    if align: x.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    return x
ws['A1'] = '月度報告 / BÁO CÁO THÁNG THEO CHUYỀN'
ws['A1'].font = Font(name=F0, size=14, bold=True, color='FFFFFFFF'); ws['A1'].fill = F_TITLE
ws.merge_cells('A1:P1'); ws.row_dimensions[1].height = 28
cellf(2, 1, 'CHỌN THÁNG (trống = tháng mới nhất):', font=SEC, border=False); ws.merge_cells('A2:E2')
cellf(2, 6, None, fmt='mm/yyyy', fill=F_IN)
cellf(2, 8, 'ĐANG XEM:', font=SEC, border=False); ws.merge_cells('H2:I2')
cellf(2, 10, f'=IF(F2<>"",DATE(YEAR(F2),MONTH(F2),1),DATE(YEAR(MAX({R("A")})),MONTH(MAX({R("A")})),1))', fmt='mm/yyyy',
      font=Font(name=F0, size=12, bold=True, color='FFC00000'), fill=F_FX)
ws.column_dimensions['A'].width = 16
for c in range(2, 33): ws.column_dimensions[L(c)].width = 6.5
for c in range(33, 36): ws.column_dimensions[L(c)].width = 13
LINES = list(range(5, 11))          # DANH_SACH_CHUYEN rows mirrored (6 lines)
D = lambda c: f'($J$2+{c}-2)'      # date of day-column c (B=day 1)
DAYOK = lambda c, col: f'MONTH({D(c)})=MONTH($J$2)'
MEND = 'DATE(YEAR($J$2),MONTH($J$2)+1,1)'
MCOND = f'({R("A")}>=$J$2)*({R("A")}<{MEND})'
HS = f'NHAP_LIEU!$H$5:$S${N}'
def m(line_ref, c):  # matched NHAP_LIEU row for line + day
    return f'MATCH(1,INDEX(({R("A")}={D(c)})*({R("B")}={line_ref}),0),0)'

def day_header(r, title):
    cellf(r - 1, 1, title, font=SEC, border=False)
    cellf(r, 1, 'CHUYỀN', font=HDR, fill=F_VN, align='center')
    for d in range(1, 32):
        c = d + 1
        cellf(r, c, f'=IF(MONTH($J$2+{d}-1)=MONTH($J$2),{d},"")', font=HDR, fill=F_VN, align='center')

# 1. actual output per line per day
r0 = 5; day_header(r0, '1. SẢN LƯỢNG THỰC TẾ THEO NGÀY')
for k, t in ((33, 'TỔNG THÁNG'), (34, 'MỤC TIÊU (ngày đã làm)'), (35, '% ĐẠT')):
    cellf(r0, k, t, font=HDR, fill=F_VN, align='center')
ws.row_dimensions[r0].height = 30
for i, src in enumerate(LINES):
    r = r0 + 1 + i
    cellf(r, 1, f'=IF(DANH_SACH_CHUYEN!B{src}="","",DANH_SACH_CHUYEN!B{src})', font=BOLD)
    for d in range(1, 32):
        c = d + 1; mm = m(f'$A{r}', c)
        cellf(r, c, f'=IF(OR($A{r}="",MONTH({D(c)})<>MONTH($J$2)),"",IFERROR(IF(COUNT(INDEX({HS},{mm},0))=0,"",SUM(INDEX({HS},{mm},0))),""))', fmt=INT)
    cellf(r, 33, f'=IF($A{r}="","",SUM(B{r}:AF{r}))', fmt=INT, font=BOLD)
    cellf(r, 34, f'=IF($A{r}="","",SUMPRODUCT(({R("B")}=$A{r})*{MCOND}*{ANYH}*{CAT}))', fmt=INT)
    cellf(r, 35, f'=IF(OR($A{r}="",N(AH{r})=0),"",AG{r}/AH{r})', fmt=PCT, font=BOLD)
rt = r0 + 1 + len(LINES)
cellf(rt, 1, 'TỔNG', font=BOLD, fill=F_VN)
for c in range(2, 33):
    cellf(rt, c, f'=IF(COUNT({L(c)}{r0+1}:{L(c)}{rt-1})=0,"",SUM({L(c)}{r0+1}:{L(c)}{rt-1}))', fmt=INT, font=BOLD, fill=F_VN)
cellf(rt, 33, f'=SUM(AG{r0+1}:AG{rt-1})', fmt=INT, font=BOLD, fill=F_VN)
cellf(rt, 34, f'=SUM(AH{r0+1}:AH{rt-1})', fmt=INT, font=BOLD, fill=F_VN)
cellf(rt, 35, f'=IF(N(AH{rt})=0,"",AG{rt}/AH{rt})', fmt=PCT, font=BOLD, fill=F_VN)
rtg = rt + 1
cellf(rtg, 1, 'MỤC TIÊU NGÀY', font=BOLD, fill=F_VN)
for d in range(1, 32):
    c = d + 1
    cellf(rtg, c, f'=IF(MONTH({D(c)})<>MONTH($J$2),"",IF(COUNTIF({R("A")},{D(c)})=0,"",SUMPRODUCT(({R("A")}={D(c)})*{CAT})))', fmt=INT, fill=F_VN)

# 2. % achieved per line per day
r1 = rtg + 3; day_header(r1, '2. % ĐẠT MỤC TIÊU NGÀY')
for i, src in enumerate(LINES):
    r = r1 + 1 + i; ra = r0 + 1 + i
    cellf(r, 1, f'=A{ra}', font=BOLD)
    for d in range(1, 32):
        c = d + 1; mm = m(f'$A{r}', c)
        tgt = f'ROUND(INDEX(CAU_HINH_CA!$C$5:$C$50,MATCH(INDEX({R("D")},{mm}),CAU_HINH_CA!$A$5:$A$50,0))*INDEX({R("E")},{mm}),0)'
        cellf(r, c, f'=IF({L(c)}{ra}="","",IFERROR({L(c)}{ra}/{tgt},""))', fmt='0%')
rows_pct = (r1 + 1, r1 + len(LINES))
ws.conditional_formatting.add(f'B{rows_pct[0]}:AF{rows_pct[1]}', FormulaRule(formula=[f'AND(ISNUMBER(B{rows_pct[0]}),B{rows_pct[0]}<1)'], fill=PatternFill('solid', fgColor='FFF8CBAD')))
ws.conditional_formatting.add(f'B{rows_pct[0]}:AF{rows_pct[1]}', FormulaRule(formula=[f'AND(ISNUMBER(B{rows_pct[0]}),B{rows_pct[0]}>=1)'], fill=PatternFill('solid', fgColor='FFC6EFCE')))

# 3. reasons
r2 = rows_pct[1] + 3
cellf(r2 - 1, 1, '3. LÝ DO KHÔNG ĐẠT TRONG THÁNG', font=SEC, border=False)
hdr_cols = [1, 4, 7, 10]
for k, t in enumerate(['LÝ DO', 'SỐ LẦN', 'PHÚT DỪNG MÁY', 'SẢN LƯỢNG THIẾU']):
    c0 = hdr_cols[k]; cellf(r2, c0, t, font=HDR, fill=F_VN, align='center'); ws.merge_cells(start_row=r2, start_column=c0, end_row=r2, end_column=c0 + 2)
NRS = 8
for i in range(NRS):
    r = r2 + 1 + i; src = 5 + i
    vals = [f'=IF(DANH_SACH_LY_DO!A{src}="","",DANH_SACH_LY_DO!A{src})',
            f'=IF(A{r}="","",SUMPRODUCT(({R("AI")}=A{r})*{MCOND}))',
            f'=IF(A{r}="","",SUMPRODUCT(({R("AI")}=A{r})*{MCOND}*{R("AJ")}))',
            f'=IF(A{r}="","",SUMPRODUCT(({R("AI")}=A{r})*{MCOND}*{CAT})-SUMPRODUCT(({R("AI")}=A{r})*{MCOND}*{HS}))']
    for k, v in enumerate(vals):
        c0 = hdr_cols[k]; cellf(r, c0, v, fmt=None if k == 0 else INT, font=BOLD if k == 0 else None)
        ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 2)

# 4. productivity & defects per line
r3 = r2 + NRS + 3
cellf(r3 - 1, 1, '4. NĂNG SUẤT & HÀNG LỖI THEO CHUYỀN', font=SEC, border=False)
p_cols = [1, 4, 7, 10, 13, 16]
for k, t in enumerate(['CHUYỀN', 'TỔNG SẢN LƯỢNG', 'CÔNG-GIỜ', 'SP / NGƯỜI / GIỜ', 'SỐ LỖI', 'TỶ LỆ LỖI']):
    c0 = p_cols[k]; cellf(r3, c0, t, font=HDR, fill=F_VN, align='center'); ws.merge_cells(start_row=r3, start_column=c0, end_row=r3, end_column=c0 + 2)
for i in range(len(LINES)):
    r = r3 + 1 + i; ra = r0 + 1 + i
    LC = f'({R("B")}=A{r})*{MCOND}'
    vals = [(f'=A{ra}', None), (f'=IF(A{r}="","",AG{ra})', INT),
            (f'=IF(A{r}="","",SUMPRODUCT({LC}*{R("AG")}*{NHRS}))', INT),
            (f'=IF(OR(A{r}="",N(G{r})=0),"",D{r}/G{r})', '0.0'),
            (f'=IF(A{r}="","",SUMIFS({HLR("F")},{HLR("C")},A{r},{HLR("A")},">="&$J$2,{HLR("A")},"<"&{MEND}))', INT),
            (f'=IF(OR(A{r}="",N(D{r})=0),"",M{r}/D{r})', '0.00%')]
    for k, (v, f) in enumerate(vals):
        c0 = p_cols[k]; cellf(r, c0, v, fmt=f, font=BOLD if k == 0 else None)
        ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 2)

# chart: daily total actual (bars) vs daily target (line)
bar = BarChart(); bar.type = 'col'; bar.title = 'Sản lượng thực tế và mục tiêu theo ngày'
bar.y_axis.title = 'PCS'; bar.x_axis.title = 'Ngày'
bar.add_data(Reference(ws, min_col=1, max_col=32, min_row=rt), titles_from_data=True, from_rows=True)
bar.set_categories(Reference(ws, min_col=2, max_col=32, min_row=r0))
ln = LineChart(); ln.add_data(Reference(ws, min_col=1, max_col=32, min_row=rtg), titles_from_data=True, from_rows=True)
bar += ln
bar.height = 9; bar.width = 30
ws.add_chart(bar, f'A{r3 + len(LINES) + 3}')
ws.freeze_panes = 'B4'

# ---------------- HUONG_DAN ----------------
ws = wb.create_sheet('HUONG_DAN', 0)
ws.column_dimensions['A'].width = 4; ws.column_dimensions['B'].width = 110
lines = [
 ('HƯỚNG DẪN SỬ DỤNG – THEO DÕI SẢN LƯỢNG V20', 'title'),
 ('Mỗi sheet dữ liệu là một Excel Table (lọc/sắp xếp ở dòng 4). Cột có tiêu đề VÀNG (dòng 3) là cột nhập tay;', ''),
 ('cột có tiêu đề XANH NHẠT là công thức, KHÔNG gõ đè. Thêm dữ liệu: gõ vào dòng ngay dưới bảng, bảng tự nới ra và tự điền công thức.', ''),
 ('Dữ liệu bắt đầu từ dòng 5 ở mọi sheet. Dòng 2–3 là tiêu đề Trung/Việt, dòng 4 là số thứ tự cột.', ''),
 ('', ''),
 ('1. Danh mục (ít khi đổi)', 'h'),
 ('CAU_HINH_CA: mã ca và các đợt làm việc. TỔNG GIỜ tự tính từ giờ bắt đầu/kết thúc.', ''),
 ('DANH_SACH_CHUYEN: thứ tự ở đây là thứ tự hiển thị trên TV. Thêm chuyền thì thêm 1 dòng ở cả sheet này và bảng HIEN_THI.', ''),
 ('DANH_SACH_SAN_PHAM: mọi mã sản phẩm. Mã được lưu dạng chữ (883 và "883" là một).', ''),
 ('', ''),
 ('2. Đầu tháng', 'h'),
 ('MUC_TIEU_THANG: mỗi sản phẩm một dòng cho mỗi tháng. Cột THÁNG gõ dạng 10/2026. Các cột lũy kế tự tính.', ''),
 ('', ''),
 ('3. Hằng ngày – sheet NHAP_LIEU', 'h'),
 ('Mỗi chuyền một dòng mỗi ngày: NGÀY, CHUYỀN, MÃ SẢN PHẨM, MÃ CA, MỤC TIÊU MỖI GIỜ.', ''),
 ('Sau mỗi giờ, gõ số lượng làm được vào cột GIỜ 1 … GIỜ 12 (GIỜ 1 = đợt đầu tiên của ca).', ''),
 ('Máy tự tính: mục tiêu ngày = giờ ca × mục tiêu mỗi giờ (làm tròn), thực tế, % đạt, còn thiếu, tiến độ theo giờ,', ''),
 ('và THIẾU HÔM TRƯỚC = số còn thiếu của ngày làm việc gần nhất trước đó của cùng chuyền.', ''),
 ('Cột KIỂM TRA phải là OK. Nếu báo TRÙNG / THIẾU MÃ CA / CHƯA CÓ MỤC TIÊU THÁNG thì sửa dòng đó.', ''),
 ('', ''),
 ('4. Màn hình TV – sheet HIEN_THI', 'h'),
 ('Toàn bộ là công thức, KHÔNG nhập tay. App TV chỉ đọc sheet này: mỗi dòng của bảng là một chuyền, dòng TỔNG CỘNG nằm ngay dưới bảng (cách 1 dòng).', ''),
 ('Mặc định hiển thị NGÀY MỚI NHẤT có trong NHAP_LIEU. Muốn xem ngày khác thì gõ ngày vào ô J1; phải xóa J1 thì bảng mới tự theo ngày mới.', ''),
 ('Không đổi tên sheet, không chèn/xóa cột ở HIEN_THI, vì app đọc theo vị trí cột.', ''),
 ('', ''),
 ('5. Lịch làm việc, lý do, công nhân, hàng lỗi (mới ở V19)', 'h'),
 ('LICH_LAM_VIEC: Chủ nhật mặc định nghỉ. Ngày thường được nghỉ thì thêm dòng LOẠI = Nghỉ; Chủ nhật đi làm thì LOẠI = Làm bù.', ''),
 ('Từ lịch này file tính số ngày làm việc còn lại và CẦN LÀM MỖI NGÀY ĐỂ KỊP THÁNG = (mục tiêu tháng − lũy kế trước ngày đang xem) / ngày còn lại.', ''),
 ('NHAP_LIEU có thêm SỐ CÔNG NHÂN, LÝ DO KHÔNG ĐẠT (chọn từ DANH_SACH_LY_DO), PHÚT DỪNG MÁY. Ngày không đạt mà chưa ghi lý do thì KIỂM TRA sẽ báo.', ''),
 ('MUC_TIEU_THANG có thêm THIẾU THÁNG TRƯỚC (chỉ hiển thị, không cộng vào mục tiêu), số ngày làm việc và bình quân cần mỗi ngày.', ''),
 ('BAO_CAO_THANG: chọn tháng ở F2 (trống = tháng mới nhất). Có sản lượng và % đạt theo chuyền × ngày, tổng hợp lý do, năng suất, hàng lỗi và biểu đồ.', ''),
 ('', ''),
 ('6. Hàng lỗi kèm ảnh – sheet HANG_LOI (mới ở V20)', 'h'),
 ('Mỗi lần QC phát hiện hàng lỗi, thêm 1 dòng: NGÀY, GIỜ, CHUYỀN, LOẠI LỖI (chọn từ DANH_SACH_LOAI_LOI), SỐ LƯỢNG, TÊN FILE ẢNH.', ''),
 ('Chép ảnh vào thư mục images\\hang_loi nằm cạnh file Excel này, rồi gõ đúng tên file (ví dụ rach_chuyen2.jpg) vào cột TÊN FILE ẢNH.', ''),
 ('MÃ SẢN PHẨM tự lấy theo NHAP_LIEU. SỐ LỖI ở NHAP_LIEU và HIEN_THI tự cộng từ sheet này, không gõ tay nữa.', ''),
 ('Cột HIỆN TRÊN TV = CÓ là dòng của ngày đang hiển thị: TV chỉ chiếu ảnh của các dòng này. Cột KIỂM TRA phải là OK.', ''),
 ('', ''),
 ('7. Dung lượng', 'h'),
 ('Công thức tổng hợp tính tới dòng 10.004 của NHAP_LIEU (khoảng 4–5 năm với 6 chuyền).', ''),
]
for i, (t, k) in enumerate(lines, 1):
    c = ws.cell(i, 2, t)
    c.font = Font(name=F0, size=16 if k == 'title' else 12, bold=k in ('title', 'h'), color='FF1F4E78' if k else 'FF000000')
    c.alignment = Alignment(wrap_text=True, vertical='top')
ws['B44'] = 'Tiêu đề cột nhập tay'; ws['B44'].fill = F_HIN
ws['B45'] = 'Tiêu đề cột công thức'; ws['B45'].fill = F_VN

wb.active = wb.sheetnames.index('NHAP_LIEU')
wb.calculation.fullCalcOnLoad = True
out = sys.argv[1]
wb.save(out)
print('saved', out)
